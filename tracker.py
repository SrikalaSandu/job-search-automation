"""
Appends a new row to the Applications tab of Job_Application_Tracker.xlsx
(the same file structure delivered earlier). Creates the file from scratch
with the right headers/formatting if it doesn't exist yet.
"""

import os
from datetime import date
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from config import TRACKER_PATH

HEADERS = [
    "Company", "Role", "Job Board / Source", "Location", "Match %",
    "Status", "Date Found", "Date Applied", "Resume Version",
    "Cover Letter Version", "Follow-up Date", "Next Action", "Link", "Notes"
]


def _create_new_tracker(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applications"

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        c = ws.cell(row=1, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.freeze_panes = "A2"

    widths = [18, 26, 20, 18, 10, 14, 12, 12, 22, 22, 14, 28, 26, 30]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    dv_status = DataValidation(
        type="list",
        formula1='"Not Applied,Applying,Applied,OA/Assessment,Phone Screen,Interview,Offer,Rejected,Withdrawn,Ghosted"',
        allow_blank=True,
    )
    ws.add_data_validation(dv_status)
    dv_status.add("F2:F1000")

    wb.save(path)
    return wb


def add_job_to_tracker(job: dict, score_result: dict, doc_paths: Optional[dict] = None):
    """job: normalized job dict. score_result: output of matcher.score_job.
    doc_paths: output of resume_writer.generate_application_docs, or None
    if the job scored below the auto-tailor threshold."""

    os.makedirs(os.path.dirname(TRACKER_PATH), exist_ok=True)

    if os.path.exists(TRACKER_PATH):
        wb = openpyxl.load_workbook(TRACKER_PATH)
        ws = wb["Applications"]
    else:
        wb = _create_new_tracker(TRACKER_PATH)
        ws = wb["Applications"]

    next_row = ws.max_row + 1

    resume_version = os.path.basename(doc_paths["resume_path"]) if doc_paths else ""
    cover_letter_version = os.path.basename(doc_paths["cover_letter_path"]) if doc_paths else ""

    row_values = [
        job.get("company", ""),
        job.get("title", ""),
        job.get("source", ""),
        job.get("location", ""),
        (score_result.get("match_score", 0) or 0) / 100,  # stored as fraction for % format
        "Not Applied",
        date.today().isoformat(),
        "",
        resume_version,
        cover_letter_version,
        "",
        "; ".join(score_result.get("eligibility_flags", [])) or "Review and apply",
        job.get("url", ""),
        score_result.get("reasoning", ""),
    ]

    for col, val in enumerate(row_values, 1):
        cell = ws.cell(row=next_row, column=col, value=val)
        if col == 5:
            cell.number_format = "0%"

    wb.save(TRACKER_PATH)
    print(f"[tracker] Added row {next_row}: {job.get('company')} — {job.get('title')}")
